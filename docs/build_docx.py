# docs/build_docx.py — regenerate the Word version of GUIDE.md.
# Usage (from the project folder):  python docs/build_docx.py
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import GUIDE_DOCX_PATH, GUIDE_PATH  # noqa: E402

ACCENT = RGBColor(0xD7, 0x26, 0x3D)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def shade(paragraph, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = ACCENT
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_code(doc, lines):
    for line in lines or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(12)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        shade(p, "F1F5F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j, value in enumerate(row):
            paragraph = table.cell(i, j).paragraphs[0]
            add_inline(paragraph, value)
            if i == 0:
                for run in paragraph.runs:
                    run.bold = True
    doc.add_paragraph()


def build(markdown):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = markdown.splitlines()
    i, paragraph_lines = 0, []

    def flush():
        if paragraph_lines:
            add_inline(doc.add_paragraph(), " ".join(paragraph_lines))
            paragraph_lines.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
        elif stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith("# "):
            flush()
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            flush()
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            flush()
            doc.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^\s*[-*] ", line):
            flush()
            p = doc.add_paragraph(style="List Bullet 2" if line.startswith("  ") else "List Bullet")
            add_inline(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\d+\. ", stripped):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-14)
            add_inline(p, stripped)
        elif stripped.startswith(">"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            shade(p, "FEF2F2")
            add_inline(p, stripped.lstrip("> "))
        elif not stripped:
            flush()
        else:
            paragraph_lines.append(stripped)
        i += 1
    flush()
    return doc


if __name__ == "__main__":
    build(GUIDE_PATH.read_text(encoding="utf-8")).save(GUIDE_DOCX_PATH)
    print(f"✅ Wrote {GUIDE_DOCX_PATH}")
